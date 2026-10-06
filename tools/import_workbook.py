#!/usr/bin/env python3
"""One-time import of the old scheduled task's data into the SCOPE Scout app.

Reads (never writes) the old scope_scout folder:
  SCOPE_coop_picks_latest.xlsx   picks, near misses, in progress, RBC and all postings
  raw/*.json                     full posting text saved by past runs
  reader_brief.md                the candidate profile the scoring prompt uses
and writes the app's store file (default %APPDATA%\\SCOPE Scout\\scope-scout.json).

Close the app first. An existing store is only replaced with --force.

  python tools/import_workbook.py [--src DIR] [--out FILE] [--force]

Every posting that was already read gets a score (a partial one when the workbook kept
no details), so the app never re-reads or re-bills old postings.
"""
import argparse
import datetime as dt
import glob
import json
import os
import re
import sys

from openpyxl import load_workbook

STATUSES = {'To apply', 'Drafting', 'Applied', 'Skip'}


def iso_utc(d):
    """Naive local datetime from the workbook -> the app's ISO UTC timestamp format."""
    if not isinstance(d, dt.datetime):
        return None
    return d.astimezone(dt.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.000Z')


def deadline_fields(d):
    """Workbook deadline (naive Pacific time) -> (SCOPE-style text, wall-clock ISO)."""
    if not isinstance(d, dt.datetime):
        return '', None
    return f'{d:%b} {d.day}, {d.year} {d:%I:%M %p}', d.strftime('%Y-%m-%dT%H:%M')


def job_id(v):
    if v in (None, ''):
        return None
    return str(int(v)) if isinstance(v, (int, float)) else str(v).strip()


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def text(v):
    return '' if v is None else str(v).strip()


def rows_of(wb, name):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        pid = job_id(d.get('Job ID'))
        if pid:
            d['Job ID'] = pid
            out.append(d)
    return out


def split_missing(s):
    """'Required: SQL; C# | Preferred: AWS' -> (['SQL', 'C#'], ['AWS'])."""
    req, pref = [], []
    for part in text(s).split(' | '):
        if part.startswith('Required: '):
            body = part[len('Required: '):]
            if body != 'none missing':
                req = [x.strip() for x in body.split('; ') if x.strip()]
        elif part.startswith('Preferred: '):
            pref = [x.strip() for x in part[len('Preferred: '):].split('; ') if x.strip()]
    return req, pref


def full_score(d, section):
    req_missing, pref_missing = split_missing(d.get('Missing requirements'))
    s = {
        'section': section,
        'fit': num(d.get('Fit')),
        'requiredMet': num(d.get('Req met')),
        'requiredTotal': num(d.get('Req total')),
        'preferredMet': num(d.get('Pref met')),
        'preferredTotal': num(d.get('Pref total')),
        'thinDescription': d.get('Thin') == 'Yes',
        'requiredMissing': req_missing,
        'preferredMissing': pref_missing,
        'whyItFits': text(d.get('Why it fits')),
        'gaps': text(d.get('Gaps / flags')),
        'specialInstructions': text(d.get('Special instructions')),
        'city': text(d.get('City')),
        'workMode': text(d.get('Work mode')),
        'duration': text(d.get('Duration')),
        'applyVia': text(d.get('Apply via')),
        'coverLetter': text(d.get('Cover letter')),
        'salary': text(d.get('Salary')),
        'scoredAt': iso_utc(d.get('Added')),
        'model': 'imported',
    }
    return {k: v for k, v in s.items() if v is not None}


def extract_profile(path):
    """The candidate-specific parts of reader_brief.md; the scoring rules live in the app."""
    md = open(path, encoding='utf-8').read()
    parts = []
    m = re.search(r'The run date is given in your prompt\.\s*(.+)', md)
    if m:
        parts.append(m.group(1).strip())
    m = re.search(r'(## Candidate evidence.*?)(?=\n## )', md, re.S)
    if m:
        parts.append(m.group(1).strip())
    m = re.search(r'^Location notes:.*$', md, re.M)
    if m:
        parts.append('## Locations\n\n' + m.group(0).strip())
    return '\n\n'.join(parts)


def load_raw(src):
    """Full posting fields from every raw/*.json dump; newer files win."""
    details = {}
    files = sorted(glob.glob(os.path.join(src, 'raw', '*.json')),
                   key=lambda f: re.search(r'(\d{4}-\d{2}-\d{2})', f).group(1) if re.search(r'\d{4}-\d{2}-\d{2}', f) else '')
    for f in files:
        d = json.load(open(f, encoding='utf-8'))
        postings = d.get('postings', []) if isinstance(d, dict) else d
        m = re.search(r'(\d{4}-\d{2}-\d{2})', os.path.basename(f))
        fetched = iso_utc(dt.datetime.strptime(m.group(1), '%Y-%m-%d')) if m else None
        for p in postings:
            pid = job_id(p.get('postingId'))
            if pid:
                details[pid] = ({k: text(v) for k, v in p.items() if k != 'postingId'}, fetched)
    return details


def term_from_details(fields, fallback):
    term = fields.get('Placement Term', '')
    return 'S27' if 'Summer' in term else 'W27' if 'Winter' in term else fallback


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--src', default=os.path.join(os.path.expanduser('~'), 'Downloads', 'scope_scout'))
    ap.add_argument('--out', default=os.path.join(os.environ.get('APPDATA', ''), 'SCOPE Scout', 'scope-scout.json'))
    ap.add_argument('--force', action='store_true', help='replace an existing store')
    a = ap.parse_args()
    if os.path.exists(a.out) and not a.force:
        sys.exit(f'{a.out} already exists. Close the app and rerun with --force to replace it.')

    wb = load_workbook(os.path.join(a.src, 'SCOPE_coop_picks_latest.xlsx'))
    now = dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.000Z')
    postings = {}

    def ensure(pid, title='', org='', location='', deadline=None, term='W27', first_seen=None):
        if pid not in postings:
            dtext, diso = deadline_fields(deadline)
            postings[pid] = {
                'id': pid, 'term': term,
                'listing': {'title': text(title), 'org': text(org), 'location': text(location),
                            'deadlineText': dtext, 'deadline': diso, 'applicants': None, 'appStatus': '-'},
                'firstSeen': first_seen or now, 'lastSeen': first_seen or now,
                'onScopeNow': True, 'myStatus': '',
            }
        return postings[pid]

    # Every posting ever listed. "Read in full?" decides whether triage already happened.
    for d in rows_of(wb, 'All postings'):
        p = ensure(d['Job ID'], d.get('Title'), d.get('Organization'), d.get('Location'),
                   d.get('Deadline (Pacific)'), text(d.get('Term')) or 'W27', iso_utc(d.get('First seen')))
        p['triage'] = 'read' if d.get('Read in full?') == 'Yes' else 'skip'

    def set_status(p, value):
        if text(value) in STATUSES and not p['myStatus']:
            p['myStatus'] = text(value)

    for tab, term in (('January picks', 'W27'), ('May backups', 'S27')):
        for d in rows_of(wb, tab):
            p = ensure(d['Job ID'], d.get('Title'), d.get('Organization'), d.get('City'), d.get('Deadline (Pacific)'), term)
            p['score'] = full_score(d, 'pick')
            p['triage'] = 'read'
            p['listing']['applicants'] = num(d.get('Applicants so far'))
            set_status(p, d.get('Status'))

    for d in rows_of(wb, 'Near misses'):
        p = ensure(d['Job ID'], d.get('Title'), d.get('Organization'), d.get('City'))
        score = {'section': 'near', 'fit': num(d.get('Fit')), 'whyMissed': text(d.get('Why it missed')),
                 'city': text(d.get('City')), 'scoredAt': iso_utc(d.get('Added')), 'model': 'imported'}
        p['score'] = {k: v for k, v in score.items() if v is not None}
        p['triage'] = 'read'

    # RBC rows carry full scores for postings that sat on no other tab.
    for d in rows_of(wb, 'RBC'):
        p = ensure(d['Job ID'], d.get('Title'), 'RBC', d.get('City'), d.get('Deadline (Pacific)'), text(d.get('Term')) or 'W27')
        if num(d.get('Fit')) is not None and 'score' not in p:
            p['score'] = full_score(d, 'none')
            p['triage'] = 'read'
        if num(d.get('Applicants so far')) is not None:
            p['listing']['applicants'] = num(d.get('Applicants so far'))
        set_status(p, d.get('Status'))

    # In progress: the app keeps these in progress through your status.
    for d in rows_of(wb, 'Already in progress'):
        p = ensure(d['Job ID'], d.get('Title'), d.get('Organization'), '', d.get('Deadline (Pacific)'))
        p['myStatus'] = 'Drafting' if 'draft' in text(d.get('App status')).lower() else 'Applied'
        if 'score' not in p:
            p['score'] = {'section': 'pick', 'model': 'imported'}
        p['triage'] = 'read'

    # Read in full but kept off every tab: the old run judged them not a fit.
    for p in postings.values():
        if p.get('triage') == 'read' and 'score' not in p:
            p['score'] = {'section': 'none', 'model': 'imported'}

    raw = load_raw(a.src)
    with_text = 0
    for pid, p in postings.items():
        if pid in raw:
            fields, fetched = raw[pid]
            p['details'] = fields
            if fetched:
                p['fetchedAt'] = fetched
            p['term'] = term_from_details(fields, p['term'])
            with_text += 1

    profile = extract_profile(os.path.join(a.src, 'reader_brief.md'))
    store = {'version': 1, 'profile': profile, 'postings': postings, 'runs': []}
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    tmp = a.out + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(store, f, ensure_ascii=False, indent=1)
    os.replace(tmp, a.out)

    def section(p):
        if p['myStatus'] in ('Drafting', 'Applied') or p['listing']['appStatus'] != '-':
            return 'in progress'
        return {'pick': 'pick ' + p['term'], 'near': 'near miss', 'none': 'not a fit'}.get(
            p.get('score', {}).get('section'), 'unread')

    counts = {}
    for p in postings.values():
        counts[section(p)] = counts.get(section(p), 0) + 1
    print(f'Wrote {a.out}')
    print(f'{len(postings)} postings, {with_text} with full text, profile {len(profile)} characters')
    for k in sorted(counts):
        print(f'  {k}: {counts[k]}')


if __name__ == '__main__':
    main()
